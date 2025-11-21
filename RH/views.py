import os
import uuid
import boto3
import locale
import pandas as pd
import openpyxl
from datetime import date, timedelta

from django.shortcuts import render, redirect, get_object_or_404
from django.db.models import Count, Q, Avg, F
from django.http import HttpResponse, HttpResponseForbidden
from django.template.loader import get_template, render_to_string
from django.views.generic import CreateView, DetailView, UpdateView, DeleteView, ListView, FormView, TemplateView
from django.urls import reverse_lazy
from django.forms import inlineformset_factory
from django.db import transaction
from django.conf import settings
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from xhtml2pdf import pisa

from botocore.exceptions import BotoCoreError, NoCredentialsError
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter 
from django.utils import timezone

from .models import (
    Empleado, Departamento, Puesto, MotivoInactivacion,
    TipoDocumentoOperador, DocumentoOperador, HistorialLaboral, Salario, Contrato, Hijo,
    TipoViaje, DivisionOperativa, TipoCarga
)
from .forms import (
    EmpleadoForm, DepartamentoForm, PuestoForm, MotivoInactivacionForm, 
    TipoDocumentoOperadorForm, DocumentoOperadorForm, HistorialLaboralForm, 
    SalarioForm, ContratoForm, HijoForm
)

# ==============================================================================
# === FUNCIONES AUXILIARES S3 ===
# ==============================================================================

def _subir_archivo_a_s3(archivo_obj, s3_ruta_relativa):
    """Sube un archivo a S3 y retorna la ruta relativa para guardar en BD."""
    try:
        s3_client = boto3.client(
            's3',
            aws_access_key_id=settings.AWS_ACCESS_KEY_ID,
            aws_secret_access_key=settings.AWS_SECRET_ACCESS_KEY,
            region_name=settings.AWS_S3_REGION_NAME
        )
        full_s3_path = f"{settings.AWS_MEDIA_LOCATION}/{s3_ruta_relativa}"
        
        archivo_obj.seek(0)
        s3_client.upload_fileobj(
            archivo_obj,
            settings.AWS_STORAGE_BUCKET_NAME,
            full_s3_path,
            ExtraArgs={'ACL': 'public-read'}
        )
        return s3_ruta_relativa 
    except Exception as e:
        print(f"Error S3: {e}")
        return None

def _eliminar_archivo_de_s3(ruta_completa_s3):
    """Elimina un archivo de S3 dado su path."""
    if not ruta_completa_s3: return
    try:
        s3_client = boto3.client(
            's3',
            aws_access_key_id=settings.AWS_ACCESS_KEY_ID,
            aws_secret_access_key=settings.AWS_SECRET_ACCESS_KEY,
            region_name=settings.AWS_S3_REGION_NAME
        )
        full_key = f"{settings.AWS_MEDIA_LOCATION}/{ruta_completa_s3}" if not ruta_completa_s3.startswith(settings.AWS_MEDIA_LOCATION) else ruta_completa_s3
        s3_client.delete_object(Bucket=settings.AWS_STORAGE_BUCKET_NAME, Key=full_key)
    except Exception:
        pass

# Mixin para permisos (Opcional)
def es_admin_rh(user):
    return user.is_authenticated and (user.is_superuser or user.groups.filter(name__in=['Administrador', 'RH_Admin', 'Recursos Humanos']).exists())

class RHAdminRequiredMixin(LoginRequiredMixin, UserPassesTestMixin):
    def test_func(self):
        return es_admin_rh(self.request.user)

# ==============================================================================
# === VISTAS PRINCIPALES (Dashboard) ===
# ==============================================================================

def inicio_rh(request):
    today = date.today()
    
    total_empleados = Empleado.objects.count()
    empleados_activos = Empleado.objects.filter(activo=True).count()
    total_departamentos = Departamento.objects.count()
    empleados_inactivos = Empleado.objects.filter(activo=False).count()

    porcentaje_activos = round((empleados_activos / total_empleados * 100), 1) if total_empleados > 0 else 0
    porcentaje_inactivos = round((empleados_inactivos / total_empleados * 100), 1) if total_empleados > 0 else 0

    # Filtramos puesto como texto (icontains)
    operadores_migmar = Empleado.objects.filter(
        activo=True, 
        empresa='MIGMAR', 
        puesto__icontains='Operador'
    ).count()
    
    operadores_marco = Empleado.objects.filter(
        activo=True, 
        empresa='MARCO_MORALES', 
        puesto__icontains='Operador'
    ).count()

    cumpleanos_hoy = []
    empleados_cumple = Empleado.objects.filter(
        fecha_nacimiento__month=today.month,
        fecha_nacimiento__day=today.day,
        activo=True
    )
    for emp in empleados_cumple:
        edad = today.year - emp.fecha_nacimiento.year
        cumpleanos_hoy.append({'empleado': emp, 'edad_a_cumplir': edad})

    alertas_rh = []
    
    # A. Contratos
    fecha_limite_contrato = today + timedelta(days=30)
    contratos_por_vencer = Contrato.objects.filter(
        tipo_contrato='DETERMINADO',
        fecha_fin__range=[today, fecha_limite_contrato],
        empleado__activo=True
    ).select_related('empleado')

    for c in contratos_por_vencer:
        dias = (c.fecha_fin - today).days
        alertas_rh.append({
            'titulo': f'Vencimiento ({dias} días)',
            'descripcion': f'{c.empleado.nombre} {c.empleado.apellido}',
            'tipo': 'warning',
            'icono': 'file-contract',
            'fecha': c.fecha_fin
        })

    # B. Documentos
    fecha_limite_docs = today + timedelta(days=15)
    docs_vencidos = DocumentoOperador.objects.filter(
        fecha_vencimiento__lte=fecha_limite_docs,
        empleado__activo=True
    ).select_related('empleado', 'tipo_documento')

    for d in docs_vencidos:
        tipo_alerta = 'danger' if d.fecha_vencimiento < today else 'warning'
        texto_dias = "VENCIDO" if d.fecha_vencimiento < today else f"Vence: {(d.fecha_vencimiento - today).days} días"
        alertas_rh.append({
            'titulo': f'{d.tipo_documento.nombre}',
            'descripcion': f'{texto_dias} - {d.empleado.nombre}',
            'tipo': tipo_alerta,
            'icono': 'id-card',
            'fecha': d.fecha_vencimiento
        })
    
    # C. Vacantes
    vacantes_activas = HistorialLaboral.objects.filter(estatus='BUSCANDO').select_related('empleado')
    for v in vacantes_activas:
        dias = (today - v.fecha_inicio).days
        puesto_nombre = v.empleado.puesto if v.empleado else "Puesto"
        alertas_rh.append({
            'titulo': 'Vacante Abierta',
            'descripcion': f'{puesto_nombre} - {dias} días',
            'tipo': 'info',
            'icono': 'user-clock',
            'fecha': v.fecha_inicio
        })

    alertas_rh.sort(key=lambda x: (x['tipo'] != 'danger', x['fecha']))

    # Gráficos
    departamento_distribucion = Empleado.objects.filter(activo=True).values('departamento__nombre').annotate(count=Count('id')).order_by('-count')
    departamento_distribucion_list = [
        {'nombre': item['departamento__nombre'] or 'Sin Asignar', 'count': item['count']} 
        for item in departamento_distribucion
    ]

    context = {
        'total_empleados': total_empleados,
        'empleados_activos': empleados_activos,
        'total_departamentos': total_departamentos,
        'empleados_inactivos': empleados_inactivos,
        'empleados_activos_porcentaje': porcentaje_activos,
        'empleados_inactivos_porcentaje': porcentaje_inactivos,
        'operadores_migmar': operadores_migmar,
        'operadores_marco': operadores_marco,
        'cumpleanos_hoy': cumpleanos_hoy,
        'departamento_distribucion': departamento_distribucion_list,
        'alertas_rh': alertas_rh,
        'today': today,
    }
    return render(request, 'rh/home.html', context)

class EmpleadoListView(LoginRequiredMixin, ListView):
    model = Empleado
    template_name = 'rh/lista_empleados.html'
    context_object_name = 'empleados'
    paginate_by = 15

    def get_queryset(self):
        queryset = super().get_queryset()
        
        nombre = self.request.GET.get('nombre', '')
        depto_id = self.request.GET.get('departamento', '')
        puesto_id = self.request.GET.get('puesto', '')
        estado = self.request.GET.get('estado', '')
        fecha_inicio = self.request.GET.get('fecha_inicio', '')
        fecha_fin = self.request.GET.get('fecha_fin', '')
        tipo_viaje_id = self.request.GET.get('tipo_viaje', '')
        empresa = self.request.GET.get('empresa', '')

        if nombre:
            queryset = queryset.filter(Q(nombre__icontains=nombre) | Q(apellido__icontains=nombre))
        if depto_id:
             queryset = queryset.filter(departamento__nombre__icontains=depto_id)
        if puesto_id:
             queryset = queryset.filter(puesto__icontains=puesto_id)
        if estado in ['0', '1']:
            queryset = queryset.filter(activo=(estado == '1'))
        if fecha_inicio:
            queryset = queryset.filter(fecha_contratacion__gte=fecha_inicio)
        if fecha_fin:
            queryset = queryset.filter(fecha_contratacion__lte=fecha_fin)
        if tipo_viaje_id:
            queryset = queryset.filter(tipo_viaje__id=tipo_viaje_id)
        if empresa:
            queryset = queryset.filter(empresa=empresa)

        sort_by = self.request.GET.get('sort', 'id')
        direction = self.request.GET.get('direction', 'desc')
        if direction == 'desc':
            if not sort_by.startswith('-'): sort_by = f'-{sort_by}'
        else:
            if sort_by.startswith('-'): sort_by = sort_by[1:]

        valid_sort_fields = ['id', 'apellido', 'puesto', 'departamento', 'fecha_contratacion', 'empresa']
        clean_sort = sort_by.replace('-', '')
        if clean_sort == 'nombre': sort_by = sort_by.replace('nombre', 'apellido')
        
        # Ajuste para ordenar por campos de texto en relaciones si es necesario
        if clean_sort == 'departamento': sort_by = sort_by.replace('departamento', 'departamento__nombre')

        queryset = queryset.order_by(sort_by) if clean_sort in valid_sort_fields else queryset.order_by('-id')
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        today = date.today()
        for empleado in context['object_list']:
            empleado.age = calculate_age(empleado.fecha_nacimiento, today)
            if empleado.fecha_contratacion:
                start_date = empleado.fecha_contratacion
                end_date = empleado.fecha_inactivacion if not empleado.activo and empleado.fecha_inactivacion else today
                empleado.dias_laborados = (end_date - start_date).days
            else:
                empleado.dias_laborados = 0
        
        context['departamentos'] = Departamento.objects.all().order_by('nombre')
        context['puestos'] = Puesto.objects.all().order_by('nombre')
        context['tipos_viaje'] = TipoViaje.objects.all().order_by('nombre')
        context['empresas'] = Empleado.EMPRESA_CHOICES
        context['total_empleados'] = self.model.objects.count()
        context['empleados_activos'] = self.model.objects.filter(activo=True).count()
        context['empleados_inactivos'] = self.model.objects.filter(activo=False).count()
        
        current_get_params = self.request.GET.copy()
        if 'page' in current_get_params: del current_get_params['page']
        context['query_string'] = current_get_params.urlencode()
        return context

def calculate_age(birth_date, current_date):
    if not birth_date: return None
    age = current_date.year - birth_date.year
    if (current_date.month, current_date.day) < (birth_date.month, birth_date.day):
        age -= 1
    return age

def cumpleanos_rh(request):
    try: locale.setlocale(locale.LC_TIME, 'es_ES.UTF-8')
    except: pass
    today = date.today()
    mes_actual = today.month
    empleados_mes = Empleado.objects.filter(fecha_nacimiento__month=mes_actual, activo=True).order_by('fecha_nacimiento__day')
    lista_cumpleanos_mes, cumpleanos_hoy = [], []

    for emp in empleados_mes:
        try: cumple_este_ano = date(today.year, emp.fecha_nacimiento.month, emp.fecha_nacimiento.day)
        except ValueError: cumple_este_ano = date(today.year, 2, 28)
        edad = today.year - emp.fecha_nacimiento.year
        dias_faltantes = (cumple_este_ano - today).days
        es_hoy = (dias_faltantes == 0)
        data = {'empleado': emp, 'dia_nacimiento': emp.fecha_nacimiento.day, 'fecha_cumple': cumple_este_ano, 'edad_a_cumplir': edad, 'dias_faltantes': dias_faltantes, 'es_hoy': es_hoy}
        lista_cumpleanos_mes.append(data)
        if es_hoy: cumpleanos_hoy.append(data)

    context = {'cumpleanos_hoy': cumpleanos_hoy, 'lista_cumpleanos_mes': lista_cumpleanos_mes, 'nombre_mes': today.strftime('%B').capitalize(), 'today': today}
    return render(request, 'rh/cumpleanos.html', context)

# ==============================================================================
# === GESTIÓN DE EMPLEADOS (CRUD + S3 + FORMSETS) ===
# ==============================================================================

DocumentoOperadorFormSet = inlineformset_factory(Empleado, DocumentoOperador, form=DocumentoOperadorForm, extra=0, can_delete=True)
HistorialLaboralFormSet = inlineformset_factory(Empleado, HistorialLaboral, form=HistorialLaboralForm, fk_name='empleado', extra=0, can_delete=True)
SalarioFormSet = inlineformset_factory(Empleado, Salario, form=SalarioForm, extra=0, can_delete=True)
ContratoFormSet = inlineformset_factory(Empleado, Contrato, form=ContratoForm, extra=0, can_delete=True)
HijoFormSet = inlineformset_factory(Empleado, Hijo, fields=['nombre', 'fecha_nacimiento'], extra=0, can_delete=True)

class EmpleadoCreateView(LoginRequiredMixin, CreateView):
    model = Empleado
    form_class = EmpleadoForm
    template_name = 'rh/empleado_form.html'
    success_url = reverse_lazy('rh:lista_empleados')

    def get_context_data(self, **kwargs):
        data = super().get_context_data(**kwargs)
        if self.request.POST:
            data['documentos_operador_formset'] = DocumentoOperadorFormSet(self.request.POST, self.request.FILES, prefix='documentos_operador')
            data['historial_laboral_formset'] = HistorialLaboralFormSet(self.request.POST, self.request.FILES, prefix='historial_laboral')
            data['salario_formset'] = SalarioFormSet(self.request.POST, prefix='salarios')
            data['contrato_formset'] = ContratoFormSet(self.request.POST, self.request.FILES, prefix='contratos')
            data['hijos_formset'] = HijoFormSet(self.request.POST, prefix='hijos')
        else:
            data['documentos_operador_formset'] = DocumentoOperadorFormSet(prefix='documentos_operador')
            data['historial_laboral_formset'] = HistorialLaboralFormSet(prefix='historial_laboral')
            data['salario_formset'] = SalarioFormSet(prefix='salarios')
            data['contrato_formset'] = ContratoFormSet(prefix='contratos')
            data['hijos_formset'] = HijoFormSet(prefix='hijos')
        return data

    def form_valid(self, form):
        context = self.get_context_data()
        formsets = {
            'documentos_operador': context['documentos_operador_formset'],
            'historial_laboral': context['historial_laboral_formset'],
            'salarios': context['salario_formset'],
            'contratos': context['contrato_formset'],
            'hijos': context['hijos_formset']
        }
        
        if all(fs.is_valid() for fs in formsets.values()):
            try:
                with transaction.atomic():
                    self.object = form.save(commit=False)
                    
                    # 1. GUARDAR PRIMERO para que la BD asigne el ID (1, 2, 3...)
                    self.object.save()
                    
                    # 2. Ahora ya tenemos ID numérico
                    carpeta_id = str(self.object.id)
                    campos_archivos = ['foto_perfil', 'ine_frente', 'ine_reverso', 'licencia']
                    
                    guardar_nuevamente = False
                    for campo in campos_archivos:
                        if hasattr(self.object, campo):
                            archivo = self.request.FILES.get(campo)
                            if archivo:
                                ext = os.path.splitext(archivo.name)[1]
                                # Ruta: rh/empleados/1/foto_perfil.jpg
                                s3_path = f"rh/empleados/{carpeta_id}/{campo}{ext}"
                                ruta_final = _subir_archivo_a_s3(archivo, s3_path)
                                if ruta_final: 
                                    setattr(self.object, campo, ruta_final)
                                    guardar_nuevamente = True

                    if guardar_nuevamente:
                        self.object.save()

                    form.save_m2m()
                    
                    for fs in formsets.values():
                        fs.instance = self.object
                        fs.save()
                
                messages.success(self.request, "Empleado registrado correctamente.")
                return redirect(self.get_success_url())
            except Exception as e:
                messages.error(self.request, f"Error al guardar: {e}")
                return self.form_invalid(form)
        else:
            return self.form_invalid(form)

    def form_invalid(self, form):
        context = self.get_context_data(form=form)
        all_errors = []
        
        if form.errors:
            for field, error_list in form.errors.items():
                if field == '__all__':
                    all_errors.append(f"Error general: {error_list[0]}")
                else:
                    field_obj = form.fields.get(field)
                    field_label = field_obj.label if field_obj else field
                    all_errors.append(f"Error en '{field_label}': {error_list[0]}")
        
        formsets = {
            'Documentos de Operador': context['documentos_operador_formset'],
            'Historial Laboral': context['historial_laboral_formset'],
            'Salarios': context['salario_formset'],
            'Contratos': context['contrato_formset'],
            'Hijos': context['hijos_formset']
        }

        for name, fs in formsets.items():
            if fs.errors:
                hay_errores = any(e for e in fs.errors if e)
                if hay_errores or fs.non_form_errors():
                    all_errors.append(f"Hay errores en la sección '{name}'. Por favor verifica los datos ingresados.")

        context['all_errors'] = all_errors
        return self.render_to_response(context)

class EmpleadoUpdateView(LoginRequiredMixin, UpdateView):
    model = Empleado
    form_class = EmpleadoForm
    template_name = 'rh/empleado_form.html'
    context_object_name = 'empleado'
    success_url = reverse_lazy('rh:lista_empleados')

    def get_context_data(self, **kwargs):
        data = super().get_context_data(**kwargs)
        if self.request.POST:
            data['documentos_operador_formset'] = DocumentoOperadorFormSet(self.request.POST, self.request.FILES, instance=self.object, prefix='documentos_operador')
            data['historial_laboral_formset'] = HistorialLaboralFormSet(self.request.POST, self.request.FILES, instance=self.object, prefix='historial_laboral')
            data['salario_formset'] = SalarioFormSet(self.request.POST, instance=self.object, prefix='salarios')
            data['contrato_formset'] = ContratoFormSet(self.request.POST, self.request.FILES, instance=self.object, prefix='contratos')
            data['hijos_formset'] = HijoFormSet(self.request.POST, instance=self.object, prefix='hijos')
        else:
            data['documentos_operador_formset'] = DocumentoOperadorFormSet(instance=self.object, prefix='documentos_operador')
            data['historial_laboral_formset'] = HistorialLaboralFormSet(instance=self.object, prefix='historial_laboral')
            data['salario_formset'] = SalarioFormSet(instance=self.object, prefix='salarios')
            data['contrato_formset'] = ContratoFormSet(instance=self.object, prefix='contratos')
            data['hijos_formset'] = HijoFormSet(instance=self.object, prefix='hijos')
        
        if self.object:
            employee = self.object
            historial = employee.historial_laboral_eventos.all()
            latest_salary = employee.salarios.order_by('-fecha_efectiva').first()
            data['dashboard_stats'] = {
                'sueldo_mensual': f"${latest_salary.sueldo_mensual:,.2f}" if latest_salary else "N/A",
                'actas_administrativas': historial.filter(tipo_evento='ACTA_ADMINISTRATIVA').count(),
            }
        return data

    def form_valid(self, form):
        context = self.get_context_data()
        formsets = {
            'documentos_operador': context['documentos_operador_formset'],
            'historial_laboral': context['historial_laboral_formset'],
            'salarios': context['salario_formset'],
            'contratos': context['contrato_formset'],
            'hijos': context['hijos_formset']
        }
        
        if all(fs.is_valid() for fs in formsets.values()):
            try:
                with transaction.atomic():
                    self.object = form.save(commit=False)
                    empleado_original = self.get_object()
                    
                    original_activo = empleado_original.activo
                    new_activo = form.cleaned_data['activo']
                    if original_activo and not new_activo:
                        self.object.fecha_inactivacion = date.today()
                    elif not original_activo and new_activo:
                        self.object.motivo_inactivacion = None
                        self.object.fecha_inactivacion = None

                    carpeta_id = str(self.object.id)
                    campos_archivos = ['foto_perfil', 'ine_frente', 'ine_reverso', 'licencia']
                    
                    for campo in campos_archivos:
                         if hasattr(self.object, campo):
                            nuevo_archivo = self.request.FILES.get(campo)
                            if nuevo_archivo:
                                archivo_viejo = getattr(empleado_original, campo)
                                if archivo_viejo: _eliminar_archivo_de_s3(archivo_viejo.name)
                                
                                ext = os.path.splitext(nuevo_archivo.name)[1]
                                s3_path = f"rh/empleados/{carpeta_id}/{campo}{ext}"
                                ruta_final = _subir_archivo_a_s3(nuevo_archivo, s3_path)
                                if ruta_final: setattr(self.object, campo, ruta_final)

                    self.object.save()
                    form.save_m2m()
                    for fs in formsets.values():
                        fs.instance = self.object
                        fs.save()
                
                messages.success(self.request, "Empleado actualizado correctamente.")
                return redirect(self.get_success_url())
            except Exception as e:
                messages.error(self.request, f"Error al actualizar: {e}")
                return self.form_invalid(form)
        else:
            return self.form_invalid(form)

    def form_invalid(self, form):
        context = self.get_context_data(form=form)
        all_errors = []
        if form.errors:
            for field, error_list in form.errors.items():
                for error in error_list:
                    field_label = form.fields.get(field).label if form.fields.get(field) else field
                    all_errors.append(f"Error en '{field_label}': {error}")
        
        formsets = {
            "Documentos de Operador": context['documentos_operador_formset'],
            "Contratos": context['contrato_formset'],
            "Historial Laboral": context['historial_laboral_formset'],
            "Salarios": context['salario_formset'],
            "Hijos": context['hijos_formset']
        }
        for name, fs in formsets.items():
            if fs.errors:
                 all_errors.append(f"Revisar errores en la sección '{name}'.")

        context['all_errors'] = all_errors
        return self.render_to_response(context)


class EmpleadoDetailView(LoginRequiredMixin, DetailView):
    model = Empleado
    template_name = 'rh/empleado_detail.html'
    context_object_name = 'empleado'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['documentos_operador'] = self.object.documentos_operador.all().select_related('tipo_documento')
        return context

class EmpleadoDeleteView(LoginRequiredMixin, DeleteView):
    model = Empleado
    template_name = 'rh/empleado_confirm_delete.html'
    context_object_name = 'empleado'
    success_url = reverse_lazy('rh:lista_empleados')


# ==============================================================================
# === IMPORTACIÓN MASIVA ROBUSTA (Auto-Increment + Catálogos) ===
# ==============================================================================

class ImportarEmpleadosExcelView(RHAdminRequiredMixin, FormView):
    template_name = 'rh/importar_empleados.html'
    success_url = reverse_lazy('rh:lista_empleados')
    
    def get(self, request, *args, **kwargs):
        return render(request, self.template_name, {'titulo': 'Migración Masiva Completa'})

    def post(self, request, *args, **kwargs):
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            messages.error(request, "Selecciona el archivo Excel.")
            return redirect(request.path)

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
            ws = wb.active
            
            creados = 0
            actualizados = 0
            errores = []
            
            # --- 1. MAPEO DINÁMICO DE ENCABEZADOS ---
            headers = {}
            for cell in ws[1]: 
                if cell.value:
                    headers[str(cell.value).strip()] = cell.column - 1

            def get_val(row_vals, col_name):
                idx = headers.get(col_name)
                if idx is not None and idx < len(row_vals):
                    val = row_vals[idx]
                    if val is None: return None
                    return str(val).strip()
                return None

            def parse_date(date_val):
                if not date_val: return None
                if hasattr(date_val, 'date'): return date_val.date()
                try:
                    clean_str = str(date_val).split(' ')[0]
                    return timezone.datetime.strptime(clean_str, '%Y-%m-%d').date()
                except:
                    return None

            with transaction.atomic():
                for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    try:
                        nombre = get_val(row, 'Nombre')
                        apellido = get_val(row, 'Apellido')

                        if not nombre or not apellido:
                            continue

                        # --- BUSQUEDA PARA EVITAR DUPLICADOS (LOGICA AUTO-INCREMENT) ---
                        curp = get_val(row, 'CURP')
                        no_empleado = get_val(row, 'No. Empleado')
                        
                        empleado = None
                        
                        # Prioridad 1: CURP
                        if curp:
                            empleado = Empleado.objects.filter(curp=curp).first()
                        # Prioridad 2: No. Empleado
                        if not empleado and no_empleado:
                            empleado = Empleado.objects.filter(numero_empleado=no_empleado).first()
                        
                        es_nuevo = False
                        if not empleado:
                            empleado = Empleado()
                            es_nuevo = True
                        
                        # --- CATÁLOGOS AUTOMÁTICOS ---
                        # Departamento
                        depto_str = get_val(row, 'Departamento')
                        if depto_str:
                            depto_obj, _ = Departamento.objects.get_or_create(
                                nombre__iexact=depto_str,
                                defaults={'nombre': depto_str, 'descripcion': 'Auto-generado por importación'}
                            )
                            empleado.departamento = depto_obj

                        # Puesto
                        puesto_str = get_val(row, 'Puesto') or "No asignado"
                        empleado.puesto = puesto_str
                        if puesto_str:
                            Puesto.objects.get_or_create(
                                nombre__iexact=puesto_str,
                                defaults={'nombre': puesto_str, 'descripcion': 'Auto-generado por importación'}
                            )

                        # Motivo Inactivación
                        motivo_str = get_val(row, 'Motivo Inactivación')
                        if motivo_str:
                            motivo_obj, _ = MotivoInactivacion.objects.get_or_create(
                                motivo__iexact=motivo_str,
                                defaults={'motivo': motivo_str}
                            )
                            empleado.motivo_inactivacion = motivo_obj

                        # --- ASIGNACIÓN DE CAMPOS ---
                        empleado.numero_empleado = no_empleado
                        empleado.nombre = nombre
                        empleado.apellido = apellido
                        empleado.fecha_contratacion = parse_date(get_val(row, 'Fecha Contratación')) or timezone.now().date()
                        empleado.fecha_nacimiento = parse_date(get_val(row, 'Fecha Nacimiento'))
                        empleado.email = get_val(row, 'Email')
                        empleado.telefono_personal = get_val(row, 'Teléfono Personal')
                        
                        # Domicilio
                        empleado.direccion = get_val(row, 'Calle y Número')
                        empleado.colonia = get_val(row, 'Colonia')
                        empleado.codigo_postal = get_val(row, 'C.P.')
                        empleado.ciudad = get_val(row, 'Ciudad')
                        empleado.estado = get_val(row, 'Estado')
                        empleado.pais = get_val(row, 'País') or 'México'
                        
                        # Legal
                        empleado.curp = curp
                        empleado.rfc = get_val(row, 'RFC')
                        empleado.nss = get_val(row, 'NSS')
                        empleado.estado_civil = get_val(row, 'Estado Civil')
                        empleado.nacionalidad = get_val(row, 'Nacionalidad')
                        
                        # Familia
                        empleado.nombre_conyuge = get_val(row, 'Nombre Cónyuge')
                        empleado.telefono_conyuge = get_val(row, 'Teléfono Cónyuge')
                        
                        # Bancario
                        empleado.banco = get_val(row, 'Banco')
                        empleado.clabe_interbancaria = get_val(row, 'CLABE')
                        empleado.numero_cuenta = get_val(row, 'No. Cuenta')
                        empleado.numero_tarjeta = get_val(row, 'No. Tarjeta')
                        
                        # Referencias
                        empleado.nombre_referencia_1 = get_val(row, 'Ref. 1 Nombre')
                        empleado.telefono_referencia_1 = get_val(row, 'Ref. 1 Tel')
                        empleado.relacion_referencia_1 = get_val(row, 'Ref. 1 Relación')
                        empleado.nombre_referencia_2 = get_val(row, 'Ref. 2 Nombre')
                        empleado.telefono_referencia_2 = get_val(row, 'Ref. 2 Tel')
                        empleado.relacion_referencia_2 = get_val(row, 'Ref. 2 Relación')

                        # Estatus
                        estatus_val = get_val(row, 'Estatus')
                        empleado.activo = True if estatus_val and 'ACTIVO' in estatus_val.upper() else False
                        
                        if not empleado.activo:
                            empleado.fecha_inactivacion = parse_date(get_val(row, 'Fecha Inactivación')) or timezone.now().date()
                        else:
                            empleado.fecha_inactivacion = None
                            empleado.motivo_inactivacion = None

                        empresa_raw = get_val(row, 'Empresa')
                        if empresa_raw:
                            if 'MARCO' in empresa_raw.upper(): empleado.empresa = 'MARCO_MORALES'
                            elif 'MIGMAR' in empresa_raw.upper(): empleado.empresa = 'MIGMAR'
                        
                        # VINCULACIÓN DE ARCHIVOS (S3)
                        S3_RUTA_FOTOS = "rh/importacion/fotos/"
                        S3_RUTA_INE = "rh/importacion/ine/"
                        foto_nombre = get_val(row, 'Nombre Archivo Foto') or get_val(row, 'Tiene Foto')
                        if foto_nombre and ('.jpg' in foto_nombre or '.png' in foto_nombre):
                             empleado.foto_perfil.name = f"{S3_RUTA_FOTOS}{foto_nombre}"

                        # GUARDAR PARA OBTENER ID
                        empleado.save()

                        # --- M2M (REQUIEREN ID PREVIO) ---
                        # Viajes
                        viajes_str = get_val(row, 'Tipos de Viaje')
                        empleado.tipo_viaje.clear()
                        if viajes_str:
                            for v in viajes_str.split(','):
                                v_limpio = v.strip()
                                if v_limpio:
                                    obj, _ = TipoViaje.objects.get_or_create(nombre__iexact=v_limpio, defaults={'nombre': v_limpio})
                                    empleado.tipo_viaje.add(obj)

                        # Carga
                        carga_str = get_val(row, 'Tipos de Carga')
                        empleado.tipo_carga.clear()
                        if carga_str:
                            for c in carga_str.split(','):
                                c_limpio = c.strip()
                                if c_limpio:
                                    obj, _ = TipoCarga.objects.get_or_create(nombre__iexact=c_limpio, defaults={'nombre': c_limpio})
                                    empleado.tipo_carga.add(obj)

                        # Divisiones
                        div_str = get_val(row, 'Divisiones Operativas')
                        empleado.division_operativa.clear()
                        if div_str:
                            for d in div_str.split(','):
                                d_limpio = d.strip()
                                if d_limpio:
                                    obj, _ = DivisionOperativa.objects.get_or_create(nombre__iexact=d_limpio, defaults={'nombre': d_limpio})
                                    empleado.division_operativa.add(obj)

                        # Salario
                        sueldo_diario = get_val(row, 'Sueldo Diario Actual')
                        if sueldo_diario:
                            try:
                                val_sueldo = float(sueldo_diario)
                                if val_sueldo > 0:
                                    ultimo = empleado.salarios.order_by('-fecha_efectiva').first()
                                    if not ultimo or abs(float(ultimo.sueldo_diario) - val_sueldo) > 0.1:
                                        Salario.objects.create(
                                            empleado=empleado,
                                            sueldo_diario=val_sueldo,
                                            fecha_efectiva=empleado.fecha_contratacion or timezone.now().date(),
                                            observaciones="Importación Masiva"
                                        )
                            except: pass

                        if es_nuevo: creados += 1
                        else: actualizados += 1
                        
                    except Exception as e:
                        errores.append(f"Fila {index} ({get_val(row, 'Nombre')}): {str(e)}")
            
            if creados > 0 or actualizados > 0: 
                messages.success(request, f"Éxito: {creados} nuevos empleados, {actualizados} actualizados.")
            if errores: 
                messages.warning(request, f"Hubo errores en {len(errores)} filas. Verifique los datos.")
                print(errores)
                
            return redirect(self.success_url)

        except Exception as e:
            messages.error(request, f"Error crítico leyendo el archivo: {e}")
            return redirect(request.path)

# ==============================================================================
# === OTRAS VISTAS CRUD ===
# ==============================================================================

class DepartamentoListView(LoginRequiredMixin, ListView):
    model = Departamento
    template_name = 'rh/lista_departamentos.html'
    context_object_name = 'departamentos'
    ordering = ['nombre']

class DepartamentoCreateView(LoginRequiredMixin, CreateView):
    model = Departamento
    form_class = DepartamentoForm
    template_name = 'rh/departamento_form.html'
    success_url = reverse_lazy('rh:lista_departamentos')

class DepartamentoDetailView(LoginRequiredMixin, DetailView):
    model = Departamento
    template_name = 'rh/departamento_detail.html'
    context_object_name = 'departamento'

class DepartamentoUpdateView(LoginRequiredMixin, UpdateView):
    model = Departamento
    form_class = DepartamentoForm
    template_name = 'rh/departamento_form.html'
    success_url = reverse_lazy('rh:lista_departamentos')

class DepartamentoDeleteView(LoginRequiredMixin, DeleteView):
    model = Departamento
    template_name = 'rh/departamento_confirm_delete.html'
    context_object_name = 'departamento'
    success_url = reverse_lazy('rh:lista_departamentos')

class PuestoListView(LoginRequiredMixin, ListView):
    model = Puesto
    template_name = 'rh/lista_puestos.html'
    context_object_name = 'puestos'
    ordering = ['nombre']

class PuestoCreateView(LoginRequiredMixin, CreateView):
    model = Puesto
    form_class = PuestoForm
    template_name = 'rh/puesto_form.html'
    success_url = reverse_lazy('rh:lista_puestos')

class PuestoDetailView(LoginRequiredMixin, DetailView):
    model = Puesto
    template_name = 'rh/puesto_detail.html'
    context_object_name = 'puesto'

class PuestoUpdateView(LoginRequiredMixin, UpdateView):
    model = Puesto
    form_class = PuestoForm
    template_name = 'rh/puesto_form.html'
    success_url = reverse_lazy('rh:lista_puestos')

class PuestoDeleteView(LoginRequiredMixin, DeleteView):
    model = Puesto
    template_name = 'rh/puesto_confirm_delete.html'
    context_object_name = 'puesto'
    success_url = reverse_lazy('rh:lista_puestos')

class MotivoInactivacionListView(LoginRequiredMixin, ListView):
    model = MotivoInactivacion
    template_name = 'rh/lista_motivos_inactivacion.html'
    context_object_name = 'motivos'
    ordering = ['motivo']

class MotivoInactivacionCreateView(LoginRequiredMixin, CreateView):
    model = MotivoInactivacion
    form_class = MotivoInactivacionForm
    template_name = 'rh/motivo_inactivacion_form.html'
    success_url = reverse_lazy('rh:lista_motivos_inactivacion')

class MotivoInactivacionUpdateView(LoginRequiredMixin, UpdateView):
    model = MotivoInactivacion
    form_class = MotivoInactivacionForm
    template_name = 'rh/motivo_inactivacion_form.html'
    success_url = reverse_lazy('rh:lista_motivos_inactivacion')

class MotivoInactivacionDeleteView(LoginRequiredMixin, DeleteView):
    model = MotivoInactivacion
    template_name = 'rh/motivo_inactivacion_confirm_delete.html'
    context_object_name = 'motivo'
    success_url = reverse_lazy('rh:lista_motivos_inactivacion')

class TipoDocumentoOperadorListView(LoginRequiredMixin, ListView):
    model = TipoDocumentoOperador
    template_name = 'rh/lista_tipos_documento_operador.html'
    context_object_name = 'tipos_documento'
    ordering = ['nombre']

class TipoDocumentoOperadorCreateView(LoginRequiredMixin, CreateView):
    model = TipoDocumentoOperador
    form_class = TipoDocumentoOperadorForm
    template_name = 'rh/tipo_documento_operador_form.html'
    success_url = reverse_lazy('rh:lista_tipos_documento_operador')

class TipoDocumentoOperadorUpdateView(LoginRequiredMixin, UpdateView):
    model = TipoDocumentoOperador
    form_class = TipoDocumentoOperadorForm
    template_name = 'rh/tipo_documento_operador_form.html'
    success_url = reverse_lazy('rh:lista_tipos_documento_operador')

class TipoDocumentoOperadorDeleteView(LoginRequiredMixin, DeleteView):
    model = TipoDocumentoOperador
    template_name = 'rh/tipo_documento_operador_confirm_delete.html'
    context_object_name = 'tipo_documento'
    success_url = reverse_lazy('rh:lista_tipos_documento_operador')

# ==============================================================================
# === REPORTES PDF / EXCEL ===
# ==============================================================================

def generar_pdf_empleado(request, pk):
    empleado = get_object_or_404(
        Empleado.objects.select_related(
             'supervisor', 'motivo_inactivacion', 'departamento'
        ).prefetch_related(
            'hijos', 'contratos', 'documentos_operador__tipo_documento',
            'salarios', 'historial_laboral_eventos', 'tipo_viaje',
            'tipo_carga', 'division_operativa'
        ),
        pk=pk
    )
    template_path = 'rh/empleado_pdf_template.html'
    context = {'empleado': empleado}
    template = get_template(template_path)
    html = template.render(context)
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="ficha-empleado-{empleado.nombre}-{empleado.apellido}.pdf"'
    pisa_status = pisa.CreatePDF(html, dest=response)
    if pisa_status.err:
       return HttpResponse('Error al generar PDF', status=500)
    return response

def semaforo_documentos_view(request):
    today = date.today()
    expirable_items = []

    documentos = DocumentoOperador.objects.filter(fecha_vencimiento__isnull=False, empleado__activo=True).select_related('empleado', 'tipo_documento')
    for doc in documentos:
        expirable_items.append({
            'empleado': doc.empleado,
            'nombre_item': f"Doc. Operador: {doc.tipo_documento.nombre}",
            'fecha_vencimiento': doc.fecha_vencimiento,
            'comentarios': doc.observaciones
        })

    contratos_determinados = Contrato.objects.filter(tipo_contrato='DETERMINADO', fecha_fin__isnull=False, empleado__activo=True).select_related('empleado')
    for contrato in contratos_determinados:
        expirable_items.append({
            'empleado': contrato.empleado,
            'nombre_item': 'Contrato Determinado',
            'fecha_vencimiento': contrato.fecha_fin,
            'comentarios': contrato.comentarios
        })

    expirable_items.sort(key=lambda x: x['fecha_vencimiento'])
    
    docs_rojo, docs_amarillo, docs_verde = [], [], []
    for item in expirable_items:
        days_remaining = (item['fecha_vencimiento'] - today).days
        item['dias_restantes'] = days_remaining
        if days_remaining <= 10: docs_rojo.append(item)
        elif days_remaining <= 30: docs_amarillo.append(item)
        else: docs_verde.append(item)

    context = {'docs_rojo': docs_rojo, 'docs_amarillo': docs_amarillo, 'docs_verde': docs_verde, 'today': today}
    return render(request, 'rh/semaforo_documentos.html', context)

def export_empleados_excel(request):
    empleados = Empleado.objects.all().select_related(
        'motivo_inactivacion', 
        'supervisor', 
        'departamento'
    ).prefetch_related(
        'division_operativa', 
        'tipo_carga', 
        'tipo_viaje',
        'salarios',
        'hijos'
    ).order_by('id')

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = 'Reporte Maestro Empleados'

    headers = [
        'ID', 'No. Empleado', 'Nombre', 'Apellido', 'Estatus', 
        'Puesto', 'Departamento', 'Empresa', 'Supervisor',
        'Fecha Contratación', 'Antigüedad (Años)', 'Fecha Inactivación', 'Motivo Inactivación',
        'Fecha Nacimiento', 'Edad', 'Email', 'Teléfono Personal', 
        'Estado Civil', 'Nacionalidad', 'CURP', 'RFC', 'NSS',
        'Calle y Número', 'Colonia', 'C.P.', 'Ciudad', 'Estado', 'País',
        'Nombre Cónyuge', 'Teléfono Cónyuge', 'Hijos (Resumen)',
        'Tipos de Viaje', 'Tipos de Carga', 'Divisiones Operativas',
        'Banco', 'CLABE', 'No. Cuenta', 'No. Tarjeta',
        'Sueldo Diario Actual', 'Sueldo Mensual Estimado',
        'Ref. 1 Nombre', 'Ref. 1 Tel', 'Ref. 1 Relación',
        'Ref. 2 Nombre', 'Ref. 2 Tel', 'Ref. 2 Relación',
        'Tiene Foto', 'Tiene INE', 'Tiene Comp. Dom.', 'Tiene CV',
        'Tiene Acta Nac.', 'Tiene Comp. Estudios', 'Tiene Const. Fiscal',
        'Tiene IMSS', 'Tiene Infonavit', 'Cartas Rec.'
    ]
    worksheet.append(headers)

    for emp in empleados:
        f_contrato = emp.fecha_contratacion.strftime('%Y-%m-%d') if emp.fecha_contratacion else ''
        f_baja = emp.fecha_inactivacion.strftime('%Y-%m-%d') if emp.fecha_inactivacion else ''
        f_nac = emp.fecha_nacimiento.strftime('%Y-%m-%d') if emp.fecha_nacimiento else ''
        
        divisiones = ", ".join([div.nombre for div in emp.division_operativa.all()])
        viajes = ", ".join([tv.nombre for tv in emp.tipo_viaje.all()])
        cargas = ", ".join([tc.nombre for tc in emp.tipo_carga.all()])
        
        ultimo_salario = emp.salarios.order_by('-fecha_efectiva').first()
        sueldo_diario = ultimo_salario.sueldo_diario if ultimo_salario else 0
        sueldo_mensual = ultimo_salario.sueldo_mensual if ultimo_salario else 0

        hijos_list = [f"{h.nombre} ({h.edad} años)" for h in emp.hijos.all()]
        hijos_str = "; ".join(hijos_list)

        def check_doc(doc_field):
            return "SÍ" if doc_field else "NO"

        cartas_count = 0
        if emp.carta_recomendacion_1_documento: cartas_count += 1
        if emp.carta_recomendacion_2_documento: cartas_count += 1

        row_data = [
            emp.id, 
            emp.numero_empleado, 
            emp.nombre, 
            emp.apellido, 
            'ACTIVO' if emp.activo else 'BAJA',
            emp.puesto, 
            emp.departamento.nombre if emp.departamento else '', 
            emp.get_empresa_display() if emp.empresa else '',
            str(emp.supervisor) if emp.supervisor else '',
            f_contrato,
            emp.antiguedad,
            f_baja,
            str(emp.motivo_inactivacion) if emp.motivo_inactivacion else '',
            f_nac,
            emp.edad,
            emp.email,
            emp.telefono_personal,
            emp.estado_civil,
            emp.nacionalidad,
            emp.curp,
            emp.rfc,
            emp.nss,
            emp.direccion,
            emp.colonia,
            emp.codigo_postal,
            emp.ciudad,
            emp.estado,
            emp.pais,
            emp.nombre_conyuge,
            emp.telefono_conyuge,
            hijos_str,
            viajes,
            cargas,
            divisiones,
            emp.banco,
            emp.clabe_interbancaria,
            emp.numero_cuenta,
            emp.numero_tarjeta,
            sueldo_diario,
            sueldo_mensual,
            emp.nombre_referencia_1, emp.telefono_referencia_1, emp.relacion_referencia_1,
            emp.nombre_referencia_2, emp.telefono_referencia_2, emp.relacion_referencia_2,
            check_doc(emp.foto_perfil),
            check_doc(emp.ine_documento),
            check_doc(emp.comprobante_domicilio),
            check_doc(emp.curriculum_vitae),
            check_doc(emp.acta_nacimiento_documento),
            check_doc(emp.comprobante_estudios_documento),
            check_doc(emp.constancia_fiscal_documento),
            check_doc(emp.semanas_cotizadas_imss_documento),
            check_doc(emp.aviso_retencion_infonavit_documento),
            f"{cartas_count} entregadas"
        ]
        worksheet.append(row_data)

    full_range = f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    tabla = Table(displayName="TablaMaestraEmpleados", ref=full_range)
    style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    tabla.tableStyleInfo = style
    worksheet.add_table(tabla)

    for col in worksheet.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value is not None:
                    cell_length = len(str(cell.value))
                    if cell_length > max_length: max_length = cell_length
            except: pass
        adjusted_width = (max_length + 2)
        if adjusted_width > 50: adjusted_width = 50
        worksheet.column_dimensions[column].width = adjusted_width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Reporte_Total_Empleados_{date.today()}.xlsx"'
    workbook.save(response)
    return response

def reporte_bajas(request):
    eventos_baja = HistorialLaboral.objects.filter(tipo_evento__in=['RENUNCIA', 'BAJA', 'ABANDONO']).select_related('empleado').order_by('-fecha_inicio')

    if request.GET.get('export') == 'excel':
        data = []
        MOTIVO_SALIDA_CHOICES_FLAT = [choice for group in HistorialLaboral.MOTIVO_SALIDA_CHOICES for choice in group[1]]
        evento_choices_dict = dict(HistorialLaboral.EVENT_CHOICES)
        motivo_choices_dict = dict(MOTIVO_SALIDA_CHOICES_FLAT)

        for evento in eventos_baja:
            data.append({
                'ID Empleado': evento.empleado.id,
                'Nombre Empleado': f"{evento.empleado.nombre} {evento.empleado.apellido}",
                'Número de Empleado': evento.empleado.numero_empleado,
                'Tipo de Evento': evento_choices_dict.get(evento.tipo_evento, evento.tipo_evento),
                'Motivo de Salida': motivo_choices_dict.get(evento.motivo_salida, evento.motivo_salida),
                'Fecha del Evento': evento.fecha_inicio,
                'Comentario': evento.descripcion,
            })
        
        df = pd.DataFrame(data)
        wb = openpyxl.Workbook()
        ws_data = wb.active
        ws_data.title = "Reporte de Bajas"
        
        for r in dataframe_to_rows(df, index=False, header=True): ws_data.append(r)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

        for col_idx, col in enumerate(ws_data.columns, 1):
            max_length = 0
            column_letter = get_column_letter(col_idx)
            for cell_idx, cell in enumerate(col):
                if cell_idx == 0:
                    cell.font = header_font
                    cell.fill = header_fill
                cell.alignment = center_align
                try:
                    if cell.value is not None:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length: max_length = cell_length
                except: pass
            ws_data.column_dimensions[column_letter].width = (max_length + 2)

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="reporte_bajas_personal.xlsx"'
        wb.save(response)
        return response

    context = {'eventos_baja': eventos_baja}
    return render(request, 'rh/reportes.html', context)

def dashboard_view(request):
    selected_company = request.GET.get('empresa', 'MIGMAR')
    
    base_operadores = Empleado.objects.filter(
        activo=True, 
        empresa=selected_company, 
        puesto__icontains='Operador'
    ).order_by('apellido', 'nombre')

    operadores_local_foraneo_list = base_operadores.filter(tipo_viaje__nombre='Local').filter(tipo_viaje__nombre='Foraneo').distinct()
    ids_ambos = set(operadores_local_foraneo_list.values_list('id', flat=True))
    operadores_locales_list = base_operadores.filter(tipo_viaje__nombre='Local').exclude(id__in=ids_ambos)
    operadores_foraneos_list = base_operadores.filter(tipo_viaje__nombre='Foraneo').exclude(id__in=ids_ambos)

    frecuencia_por_grupo = []
    for grupo in ['Autozone', 'Walmart', 'Bafar', 'Femsa', 'Sams', 'Bodega Aurrera']:
        frecuencia_por_grupo.append({'grupo': grupo, 'total': base_operadores.filter(division_operativa__nombre=grupo).count()})

    ops_autozone = set(base_operadores.filter(division_operativa__nombre='Autozone').values_list('id', flat=True))
    ops_walmart = set(base_operadores.filter(division_operativa__nombre='Walmart').values_list('id', flat=True))
    ops_sams = set(base_operadores.filter(division_operativa__nombre='Sams').values_list('id', flat=True))

    venn_data = {
        'autozone_total': len(ops_autozone),
        'walmart_total': len(ops_walmart),
        'sams_total': len(ops_sams),
        'autozone_walmart': len(ops_autozone.intersection(ops_walmart)),
        'autozone_sams': len(ops_autozone.intersection(ops_sams)),
        'walmart_sams': len(ops_walmart.intersection(ops_sams)),
        'autozone_walmart_sams': len(ops_autozone.intersection(ops_walmart).intersection(ops_sams)),
    }

    context = {
        'page_title': 'Dashboard de Operadores',
        'selected_company': selected_company,
        'total_operadores': base_operadores.count(),
        'operadores_locales': operadores_locales_list.count(),
        'operadores_foraneos': operadores_foraneos_list.count(),
        'operadores_locales_list': operadores_locales_list,
        'operadores_foraneos_list': operadores_foraneos_list,
        'operadores_local_foraneo_list': operadores_local_foraneo_list,
        'operadores_local_foraneo': operadores_local_foraneo_list.count(),
        'frecuencia_por_grupo': frecuencia_por_grupo,
        'venn_data': venn_data,
    }
    return render(request, 'rh/dashboard.html', context)

def vacantes_dashboard_view(request):
    today = date.today()
    tipos_baja = ['RENUNCIA', 'BAJA', 'ABANDONO']
    vacantes = HistorialLaboral.objects.filter(tipo_evento__in=tipos_baja).select_related('reemplazo').order_by('estatus', '-fecha_inicio')

    for vacante in vacantes:
        if vacante.estatus == 'BUSCANDO':
            vacante.dias_transcurridos = (today - vacante.fecha_inicio).days
            if vacante.empleado and vacante.empleado.puesto and vacante.empleado.departamento:
                # Como departamento ahora es FK, usamos departamento__nombre__icontains
                depto_nombre = vacante.empleado.departamento.nombre if vacante.empleado.departamento else ""
                vacante.potenciales_reemplazos = Empleado.objects.filter(
                    activo=True, 
                    puesto__icontains=str(vacante.empleado.puesto), 
                    departamento__nombre__icontains=depto_nombre
                ).exclude(id=vacante.empleado.id)
            else: vacante.potenciales_reemplazos = Empleado.objects.none()
        else:
            vacante.dias_transcurridos = (vacante.fecha_reemplazo - vacante.fecha_inicio).days if vacante.fecha_reemplazo else 0

    pie_chart_data, bar_chart_data = {}, {}
    for empresa in ['MIGMAR', 'MARCO_MORALES']:
        pie_chart_data[empresa] = {
            'activos': Empleado.objects.filter(activo=True, empresa=empresa).count(),
            'pendientes': HistorialLaboral.objects.filter(tipo_evento__in=tipos_baja, empleado__empresa=empresa, estatus='BUSCANDO').count()
        }
        avg_days_data = HistorialLaboral.objects.filter(
            empleado__empresa=empresa, estatus='REMPLAZADO', fecha_reemplazo__isnull=False
        ).values('empleado__departamento__nombre').annotate(
            avg_days=Avg(F('fecha_reemplazo') - F('fecha_inicio'))
        ).values('empleado__departamento__nombre', 'avg_days')
        
        bar_chart_data[empresa] = {
            'departamento': [
                {'name': item['empleado__departamento__nombre'], 'avg_days': item['avg_days'].days if item['avg_days'] else 0} 
                for item in avg_days_data if item['empleado__departamento__nombre']
            ]
        }

    context = {'vacantes': vacantes, 'pie_chart_data': pie_chart_data, 'bar_chart_data': bar_chart_data}
    return render(request, 'rh/vacantes_dashboard.html', context)

def asignar_reemplazo(request, pk):
    if request.method == 'POST':
        vacante = get_object_or_404(HistorialLaboral, pk=pk)
        reemplazo_id = request.POST.get('reemplazo_id')
        if reemplazo_id:
            reemplazo = get_object_or_404(Empleado, pk=reemplazo_id)
            vacante.reemplazo = reemplazo
            vacante.estatus = 'REMPLAZADO'
            vacante.fecha_reemplazo = date.today()
            vacante.save()
    return redirect('rh:vacantes_dashboard')

def reporte_documentacion_operador(request):
    todos_los_tipos_requeridos = TipoDocumentoOperador.objects.all()
    mapa_tipos_requeridos = {tipo.id: tipo.nombre for tipo in todos_los_tipos_requeridos}
    ids_tipos_requeridos = set(mapa_tipos_requeridos.keys())
    
    operadores_activos = Empleado.objects.filter(
        activo=True, 
        puesto__icontains='Operador'
    ).prefetch_related('documentos_operador')

    operadores_incompletos = []
    for operador in operadores_activos:
        ids_documentos_entregados = set(operador.documentos_operador.values_list('tipo_documento_id', flat=True))
        ids_documentos_faltantes = ids_tipos_requeridos - ids_documentos_entregados
        if ids_documentos_faltantes:
            nombres_documentos_faltantes = [mapa_tipos_requeridos[id_faltante] for id_faltante in ids_documentos_faltantes]
            operadores_incompletos.append({'operador': operador, 'documentos_faltantes': nombres_documentos_faltantes})

    context = {'operadores_incompletos': operadores_incompletos}
    return render(request, 'rh/reporte_documentacion_operador.html', context)

def descargar_plantilla_importacion(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plantilla Carga Masiva"
    
    headers = [
        'No. Empleado', 'Nombre', 'Apellido', 'Email', 
        'Fecha Contratación', 'Puesto', 'Departamento', 'Empresa',
        'Estatus', 'Fecha Inactivación', 'Motivo Inactivación',
        'Fecha Nacimiento', 'CURP', 'RFC', 'NSS',
        'Calle y Número', 'Colonia', 'C.P.', 'Ciudad', 'Estado', 'País',
        'Teléfono Personal', 'Estado Civil', 'Nacionalidad',
        'Nombre Cónyuge', 'Teléfono Cónyuge',
        'Banco', 'CLABE', 'No. Cuenta', 'No. Tarjeta',
        'Sueldo Diario Actual',
        'Tipos de Viaje', 'Tipos de Carga', 'Divisiones Operativas'
    ]
    
    ws.append(headers)
    
    # Ejemplo de llenado
    ws.append([
        '001', 'Ejemplo', 'Perez', 'ejemplo@migmar.com',
        '2024-01-01', 'Operador', 'Operaciones', 'MIGMAR',
        'ACTIVO', '', '',
        '1990-05-20', 'CURP123456...', 'RFC123...', 'NSS123...',
        'Av. Universidad 100', 'Centro', '66400', 'San Nicolas', 'NL', 'México',
        '8112345678', 'Casado/a', 'Mexicana',
        'Maria Lopez', '8187654321',
        'BBVA', '012345678901234567', '1234567890', '1234567812345678',
        '350.50',
        'Local, Foraneo', 'Seco, Refrigerado', 'Walmart, Autozone'
    ])

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except: pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="plantilla_importacion_migmar.xlsx"'
    wb.save(response)
    return response